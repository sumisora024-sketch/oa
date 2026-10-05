from __future__ import annotations

import re
import secrets
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

from app.core.config import ROOT_DIR
from app.models import BusinessPartner, ExternalDocument


STORAGE_DIR = ROOT_DIR / "backend" / "storage"
DOCUMENT_DIR = STORAGE_DIR / "external_documents"
UPLOAD_DIR = STORAGE_DIR / "external_uploads"
ASSET_DIR = STORAGE_DIR / "assets"
REFER_DIR = ROOT_DIR / "refer"

FIXED_NIT_FILES = [
    {"id": "nit-1", "name": "NIT-1 業務委託基本契約書", "filename": "NIT-1業務委託基本契約書_SES_新規NIT.docx"},
    {"id": "nit-2", "name": "NIT-2 反社会的勢力排除に関する覚書", "filename": "NIT-2反社会的勢力排除に関する覚書.pdf"},
    {"id": "nit-3", "name": "NIT-3 会社概要", "filename": "NIT-3会社概要.xlsx"},
    {"id": "nit-4", "name": "NIT-4 口座登録依頼書", "filename": "NIT-4口座登録依頼書.xlsx"},
    {"id": "nit-5", "name": "NIT-5 秘密保持契約書", "filename": "NIT-5秘密保持契約書(NDA)_210823改訂　新規取引NIT   .docx"},
]


def ensure_storage_dirs() -> None:
    for path in (DOCUMENT_DIR, UPLOAD_DIR, ASSET_DIR):
        path.mkdir(parents=True, exist_ok=True)


def fixed_file_path(file_id: str) -> Path | None:
    record = next((item for item in FIXED_NIT_FILES if item["id"] == file_id), None)
    if not record:
        return None
    path = REFER_DIR / record["filename"]
    return path if path.exists() else None


def safe_filename(value: str) -> str:
    cleaned = re.sub(r"[^\w.\-一-龥ぁ-んァ-ンー]+", "_", value, flags=re.UNICODE)
    return cleaned.strip("_") or "document"


def document_no(prefix: str) -> str:
    return f"{prefix}-{datetime.now().strftime('%Y%m%d%H%M%S')}-{secrets.token_hex(2).upper()}"


def normalize_items(items: list[dict]) -> tuple[list[dict], int, int, int]:
    normalized: list[dict] = []
    subtotal = 0
    for item in items or []:
        quantity = float(item.get("quantity") or 0)
        unit_price = int(item.get("unit_price") or 0)
        amount = item.get("amount")
        amount = int(amount) if amount is not None else round(quantity * unit_price)
        subtotal += amount
        normalized.append(
            {
                "name": item.get("name") or "SES作業費",
                "quantity": quantity,
                "unit_price": unit_price,
                "amount": amount,
                "description": item.get("description"),
            }
        )
    tax = round(subtotal * 0.1)
    total = subtotal + tax
    return normalized, subtotal, tax, total


def yen(value: int | float | None) -> str:
    if value is None:
        return "0"
    return f"{int(value):,}"


def register_pdf_font() -> str:
    font_name = "HeiseiKakuGo-W5"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont(font_name))
    return font_name


def extract_stamp_asset() -> Path | None:
    ensure_storage_dirs()
    target = ASSET_DIR / "nit_stamp.png"

    prepared = REFER_DIR / "nit_stamp.png"
    if prepared.exists():
        return prepared

    source = REFER_DIR / "请求书.png"
    if not source.exists():
        source = REFER_DIR / "发注书.png"
    if not source.exists():
        return None

    original = Image.open(source).convert("RGBA")
    image = original.copy()
    pixels = image.load()
    width, height = image.size
    mask: set[tuple[int, int]] = set()

    for y in range(height):
        for x in range(width):
            r, g, b, a = pixels[x, y]
            is_red = a > 0 and r > 145 and g < 115 and b < 115 and r > g * 1.5 and r > b * 1.5
            if is_red:
                mask.add((x, y))
            pixels[x, y] = (255, 255, 255, 0)

    if not mask:
        return None

    stamp_area = {
        (x, y)
        for x, y in mask
        if x > width * 0.78 and height * 0.28 < y < height * 0.43
    }
    if len(stamp_area) > 100:
        mask = stamp_area

    visited: set[tuple[int, int]] = set()
    filtered: set[tuple[int, int]] = set()
    for start in mask:
        if start in visited:
            continue
        stack = [start]
        component: list[tuple[int, int]] = []
        visited.add(start)
        while stack:
            x, y = stack.pop()
            component.append((x, y))
            for nx, ny in ((x + 1, y), (x - 1, y), (x, y + 1), (x, y - 1)):
                point = (nx, ny)
                if point in mask and point not in visited:
                    visited.add(point)
                    stack.append(point)
        if len(component) >= 8:
            filtered.update(component)
    if filtered:
        mask = filtered

    xs = [x for x, _ in mask]
    ys = [y for _, y in mask]
    for x, y in mask:
        pixels[x, y] = original.getpixel((x, y))

    margin = 4
    box = (
        max(min(xs) - margin, 0),
        max(min(ys) - margin, 0),
        min(max(xs) + margin, width),
        min(max(ys) + margin, height),
    )
    image.crop(box).save(target)
    return target


def draw_header(pdf: canvas.Canvas, font: str, title: str, partner: BusinessPartner, doc_no: str, issue_date: date, due_date: date | None = None) -> None:
    width, height = A4
    pdf.setFont(font, 18)
    pdf.drawCentredString(width / 2, height - 72, title)

    pdf.setFont(font, 11)
    pdf.drawString(40, height - 118, f"{partner.company_name}　御中")

    table_x = width - 205
    table_y = height - 115
    rows = [("書類番号", doc_no), ("発行日", issue_date.strftime("%Y/%m/%d"))]
    if due_date:
        rows.append(("お支払期限", due_date.strftime("%Y/%m/%d")))
    for index, (label, value) in enumerate(rows):
        y = table_y - index * 18
        pdf.rect(table_x, y, 95, 18, stroke=1, fill=0)
        pdf.rect(table_x + 95, y, 95, 18, stroke=1, fill=0)
        pdf.drawString(table_x + 5, y + 5, label)
        pdf.drawRightString(table_x + 185, y + 5, value)

    pdf.setFont(font, 12)
    pdf.drawString(width - 205, height - 250, "日本インフォテック株式会社")
    pdf.setFont(font, 9)
    pdf.drawString(width - 205, height - 270, "〒1080075")
    pdf.drawString(width - 205, height - 286, "東京都港区港南2-16-4")
    pdf.drawString(width - 205, height - 302, "品川グランドセントラルタワー8階")
    pdf.drawString(width - 205, height - 324, "Tel: 03-6863-5619")
    pdf.drawString(width - 205, height - 340, "Fax: 03-6863-4501")

    stamp = extract_stamp_asset()
    if stamp:
        pdf.drawImage(ImageReader(str(stamp)), width - 92, height - 342, width=58, height=58, mask="auto")


def draw_amount_box(pdf: canvas.Canvas, font: str, total: int) -> None:
    pdf.setFont(font, 12)
    pdf.rect(40, 500, 210, 34, stroke=1, fill=0)
    pdf.line(155, 500, 155, 534)
    pdf.drawCentredString(97, 512, "合計金額")
    pdf.drawRightString(240, 512, f"￥ {yen(total)}")


def draw_items(pdf: canvas.Canvas, font: str, items: list[dict], subtotal: int, tax: int, total: int) -> None:
    y = 455
    pdf.setFillColor(colors.lightgrey)
    pdf.rect(40, y, 515, 20, stroke=1, fill=1)
    pdf.setFillColor(colors.black)
    pdf.setFont(font, 10)
    pdf.drawString(48, y + 6, "項目")
    pdf.drawRightString(375, y + 6, "数量")
    pdf.drawRightString(450, y + 6, "単価")
    pdf.drawRightString(548, y + 6, "金額")
    y -= 20
    for item in items:
        pdf.line(40, y, 555, y)
        pdf.drawString(48, y + 6, str(item.get("name") or ""))
        pdf.drawRightString(375, y + 6, f"{item.get('quantity') or 0:g}")
        pdf.drawRightString(450, y + 6, yen(item.get("unit_price")))
        pdf.drawRightString(548, y + 6, yen(item.get("amount")))
        y -= 22
    pdf.line(40, y + 20, 555, y + 20)

    y -= 12
    pdf.drawRightString(450, y, "小計")
    pdf.drawRightString(548, y, yen(subtotal))
    y -= 20
    pdf.drawRightString(450, y, "消費税（10%）")
    pdf.drawRightString(548, y, yen(tax))
    y -= 20
    pdf.setFont(font, 11)
    pdf.drawRightString(450, y, "合計金額")
    pdf.drawRightString(548, y, yen(total))


def draw_note_box(pdf: canvas.Canvas, font: str, note: str | None, default_lines: list[str]) -> None:
    y = 150
    pdf.rect(40, 45, 515, 115, stroke=1, fill=0)
    pdf.setFont(font, 10)
    pdf.drawString(45, y, "備考")
    lines = [line for line in (note.splitlines() if note else default_lines) if line]
    for index, line in enumerate(lines[:7]):
        pdf.drawString(55, y - 16 - index * 14, line)


def create_purchase_order_pdf(path: Path, partner: BusinessPartner, document: ExternalDocument) -> None:
    font = register_pdf_font()
    pdf = canvas.Canvas(str(path), pagesize=A4)
    draw_header(pdf, font, "発注書", partner, document.document_no, document.issue_date)
    draw_amount_box(pdf, font, document.total)
    pdf.line(40, 480, 555, 480)
    draw_items(pdf, font, document.items or [], document.subtotal, document.tax, document.total)
    attrs = document.attributes or {}
    default_lines = [
        f"1.作業期間: {attrs.get('work_period', document.target_month)}",
        f"2.基本時間幅: {attrs.get('base_hours', '140時間～180時間')}　超過/控除 {attrs.get('overtime_unit', '0円')}/{attrs.get('deduction_unit', '0円')}/時間",
        f"3.契約形態: {attrs.get('contract_form', '業務委託契約')}",
        f"4.作業場所: {attrs.get('workplace', '弊社指定場所')}",
        "5.納入物件: 毎月検収当月作業報告書",
        "6.支払条件: 月末締め翌月末（振込）",
        "7.その他当注文書に記載無き事項については、別途協議と致します。",
    ]
    draw_note_box(pdf, font, document.note, default_lines)
    pdf.showPage()
    pdf.save()


def create_invoice_pdf(path: Path, partner: BusinessPartner, document: ExternalDocument) -> None:
    font = register_pdf_font()
    pdf = canvas.Canvas(str(path), pagesize=A4)
    draw_header(pdf, font, "御請求書", partner, document.document_no, document.issue_date, document.due_date)
    draw_amount_box(pdf, font, document.total)
    pdf.line(40, 480, 555, 480)
    draw_items(pdf, font, document.items or [], document.subtotal, document.tax, document.total)
    pdf.setFont(font, 10)
    bank_info = (document.attributes or {}).get("bank_info") or partner.bank_info or {}
    pdf.drawString(40, 150, "お振込先")
    pdf.drawString(40, 125, bank_info.get("bank_name", "みずほ銀行　本郷支店(075)"))
    pdf.drawString(40, 108, bank_info.get("account_type", "普通口座") + "　" + bank_info.get("account_no", "4006226"))
    pdf.drawString(40, 91, bank_info.get("account_name", "ニホンインフォテック(カ"))
    draw_note_box(pdf, font, document.note, ["恐れ入りますが、振込手数料はお客様のご負担でお願いいたします。"])
    pdf.showPage()
    pdf.save()


def create_quotation_xlsx(path: Path, partner: BusinessPartner, document: ExternalDocument) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "見積書"
    thin = Side(style="thin", color="444444")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="D9EAF7")

    ws["A1"] = "御見積書"
    ws["A1"].font = Font(size=18, bold=True)
    ws.merge_cells("A1:F1")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws["A3"] = f"{partner.company_name} 御中"
    ws["E3"] = "見積番号"
    ws["F3"] = document.document_no
    ws["E4"] = "発行日"
    ws["F4"] = document.issue_date
    ws["E5"] = "対象年月"
    ws["F5"] = document.target_month
    ws["A7"] = "合計金額"
    ws["B7"] = document.total
    ws["B7"].number_format = '"￥"#,##0'
    ws["A9"] = "項目"
    ws["D9"] = "数量"
    ws["E9"] = "単価"
    ws["F9"] = "金額"
    for cell in ws["9:9"]:
        cell.fill = header_fill
        cell.font = Font(bold=True)
        cell.border = border

    row = 10
    for item in document.items or []:
        ws.cell(row, 1, item.get("name"))
        ws.cell(row, 4, item.get("quantity"))
        ws.cell(row, 5, item.get("unit_price"))
        ws.cell(row, 6, item.get("amount"))
        ws.cell(row, 5).number_format = '"￥"#,##0'
        ws.cell(row, 6).number_format = '"￥"#,##0'
        for col in range(1, 7):
            ws.cell(row, col).border = border
        row += 1

    row += 1
    ws.cell(row, 5, "小計")
    ws.cell(row, 6, document.subtotal)
    ws.cell(row + 1, 5, "消費税(10%)")
    ws.cell(row + 1, 6, document.tax)
    ws.cell(row + 2, 5, "合計金額")
    ws.cell(row + 2, 6, document.total)
    for r in range(row, row + 3):
        ws.cell(r, 5).font = Font(bold=True)
        ws.cell(r, 6).number_format = '"￥"#,##0'

    if document.note:
        ws.cell(row + 4, 1, "備考")
        ws.cell(row + 5, 1, document.note)

    widths = [24, 14, 14, 12, 14, 16]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    wb.save(path)


def create_document_file(partner: BusinessPartner, document: ExternalDocument) -> Path:
    ensure_storage_dirs()
    ext = "xlsx" if document.document_type in {"quotation", "partner_quotation"} else "pdf"
    filename = safe_filename(f"{document.target_month}_{document.document_no}_{document.document_type}.{ext}")
    path = DOCUMENT_DIR / filename
    if document.document_type == "purchase_order":
        create_purchase_order_pdf(path, partner, document)
    elif document.document_type == "invoice":
        create_invoice_pdf(path, partner, document)
    elif document.document_type in {"quotation", "partner_quotation"}:
        create_quotation_xlsx(path, partner, document)
    else:
        raise ValueError(f"Unsupported document type: {document.document_type}")
    return path


def public_document_url(document_id: int) -> str:
    return f"/api/external/documents/{document_id}/download"
