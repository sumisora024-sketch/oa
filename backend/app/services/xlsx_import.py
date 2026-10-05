from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Employee
from app.schemas import EMPLOYEE_TYPES
from app.services.accounts import ensure_employee_user
from app.services.salary_signals import register_salary_refresh


HEADER_ALIASES = {
    "code": {"code", "员工编号", "社员编号", "编号"},
    "full_name": {"full_name", "name", "姓名", "氏名", "员工姓名"},
    "name_kana": {"name_kana", "フリガナ", "日文假名", "假名", "カナ"},
    "birth_date": {"birth_date", "出生年月日", "生日", "生年月日"},
    "graduation_status": {"graduation_status", "毕业状态", "卒業状況"},
    "residence": {"residence", "居住地", "住所"},
    "nearest_station": {"nearest_station", "最近车站", "最近的车站", "最寄駅"},
    "email": {"email", "邮箱", "メール", "mail"},
    "phone": {"phone", "电话", "手機", "電話"},
    "age": {"age", "年龄", "年齢"},
    "languages": {"languages", "语言能力", "語学"},
    "certifications": {"certifications", "证书", "資格"},
    "technical_experience": {"technical_experience", "技术经历", "技術経歴"},
    "it_years": {"it_years", "it工龄", "IT年数", "经验年数"},
    "talent_category": {"talent_category", "人才分类", "分類"},
    "skills": {"skills", "技能技术栈", "技术栈", "スキル"},
    "nationality": {"nationality", "国籍"},
    "employee_type": {"employee_type", "员工类型", "社員区分", "员工类别", "権限类型"},
}


def normalize_header(value: Any) -> str:
    text = str(value or "").strip()
    for field, aliases in HEADER_ALIASES.items():
        if text in aliases:
            return field
    return text


def parse_date(value: Any) -> date | None:
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip().replace("/", "-").replace(".", "-")
    for fmt in ("%Y-%m-%d", "%Y-%m", "%Y%m%d"):
        try:
            parsed = datetime.strptime(text, fmt)
            return parsed.date()
        except ValueError:
            continue
    return None


def parse_skills(value: Any) -> list[dict[str, str]]:
    if not value:
        return []
    if isinstance(value, list):
        return value
    items: list[dict[str, str]] = []
    for part in str(value).replace(";", ",").replace("、", ",").split(","):
        chunk = part.strip()
        if not chunk:
            continue
        if ":" in chunk:
            name, level = chunk.split(":", 1)
        elif "：" in chunk:
            name, level = chunk.split("：", 1)
        else:
            name, level = chunk, "中"
        items.append({"name": name.strip(), "level": level.strip() or "中"})
    return items


def safe_value(value: Any) -> str | None:
    text = str(value or "").strip()
    return text or None


def normalize_employee_type(value: Any) -> str:
    text = str(value or "").strip()
    return {
        "普通员工": "一般社員",
        "一般员工": "一般社員",
        "营业": "営業",
        "管理员": "管理者",
    }.get(text, text or "一般社員")


def role_for_employee_type(employee_type: str | None) -> str:
    value = (employee_type or "").strip().lower()
    if value == "hr":
        return "hr"
    if "総務" in value or "soumu" in value:
        return "soumu"
    if "管理" in value or "admin" in value:
        return "admin"
    if "営業" in value or "pm" in value or "sales" in value:
        return "pm"
    return "employee"


def normalize_graduation(value: Any) -> str | None:
    text = str(value or "").strip()
    return {
        "大学毕业": "大学卒業",
        "短大毕业": "短大卒業",
        "研究生/大学院": "大学院修了",
        "肄业或其他": "中退・その他",
    }.get(text, text or None)


def normalize_nationality(value: Any) -> str | None:
    text = str(value or "").strip()
    return {"其他": "その他"}.get(text, text or None)


def read_fixed_resume_sheet(sheet) -> dict[str, Any] | None:
    title = str(sheet["Q1"].value or sheet["A52"].value or "")
    if "技" not in title and "経歴" not in title:
        return None

    skill_pairs = [
        ("E17", "J17"), ("L17", "Q17"), ("S17", "X17"),
        ("E19", "J19"), ("L19", "Q19"), ("S19", "X19"), ("Z19", "AE19"), ("AG19", "AL19"),
        ("E25", "J25"), ("L25", "Q25"), ("S25", "X25"), ("Z25", "AE25"), ("AG25", "AL25"),
        ("E27", "J27"), ("L27", "Q27"), ("S27", "X27"), ("Z27", "AE27"), ("AG27", "AL27"), ("AN27", "AS27"),
    ]
    skills: list[dict[str, str]] = []
    for name_cell, level_cell in skill_pairs:
        name = safe_value(sheet[name_cell].value)
        if name:
            skills.append({"name": name, "level": safe_value(sheet[level_cell].value) or "中"})

    language_parts: list[str] = []
    if sheet["B46"].value:
        language_parts.append("英語: 中")
    if sheet["B47"].value:
        note = safe_value(sheet["Y47"].value)
        language_parts.append(f"日本語: {note or '中'}")

    project_texts: list[str] = []
    for row in range(58, min(sheet.max_row + 1, 120), 3):
        name = safe_value(sheet[f"G{row}"].value)
        tech = safe_value(sheet[f"V{row}"].value)
        role = safe_value(sheet[f"AG{row}"].value)
        if name:
            project_texts.append(" / ".join(part for part in [name, tech, role] if part))

    graduation = safe_value(sheet["F9"].value)
    graduate_date = parse_date(sheet["AN9"].value)
    if graduate_date:
        graduation = f"{graduation or ''} {graduate_date.isoformat()}".strip()

    return {
        "full_name": safe_value(sheet["F7"].value),
        "name_kana": safe_value(sheet["F6"].value),
        "birth_date": parse_date(sheet["R7"].value),
        "graduation_status": graduation,
        "residence": safe_value(sheet["R12"].value),
        "nearest_station": safe_value(sheet["AB13"].value),
        "email": safe_value(sheet["F13"].value),
        "phone": safe_value(sheet["F12"].value),
        "languages": " / ".join(language_parts) or None,
        "technical_experience": "\n\n".join(project_texts) or None,
        "talent_category": "infra" if any("infra" in item["name"].lower() for item in skills) else None,
        "skills": skills,
        "employee_type": "一般社員",
    }


def import_employees_from_xlsx(file_obj, db: Session) -> dict[str, Any]:
    workbook = load_workbook(file_obj, read_only=False, data_only=True)
    sheet = workbook.active
    fixed_data = read_fixed_resume_sheet(sheet)
    if fixed_data:
        return import_fixed_resume(fixed_data, db)

    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return {"created": 0, "updated": 0, "errors": ["空文件"]}

    headers = [normalize_header(value) for value in rows[0]]
    created = 0
    updated = 0
    errors: list[str] = []

    for index, row in enumerate(rows[1:], start=2):
        data = {headers[col]: row[col] for col in range(min(len(headers), len(row))) if headers[col]}
        email = str(data.get("email") or "").strip()
        full_name = str(data.get("full_name") or "").strip()
        if not email or not full_name:
            errors.append(f"第{index}行缺少姓名或邮箱")
            continue

        employee = db.scalar(select(Employee).where(Employee.email == email))
        previous_birth_date = employee.birth_date if employee else None
        was_deleted = bool(employee and employee.is_deleted)
        is_new = employee is None or was_deleted
        if is_new:
            if employee:
                employee.is_deleted = False
                employee.deleted_at = None
            else:
                employee = Employee(email=email, full_name=full_name)

        employee.code = str(data.get("code") or "").strip() or employee.code
        employee.full_name = full_name
        employee.name_kana = str(data.get("name_kana") or "").strip() or None
        employee.birth_date = parse_date(data.get("birth_date")) or employee.birth_date
        employee.graduation_status = normalize_graduation(data.get("graduation_status"))
        employee.residence = str(data.get("residence") or "").strip() or None
        employee.nearest_station = str(data.get("nearest_station") or "").strip() or None
        employee.phone = str(data.get("phone") or "").strip() or None
        employee.age = int(data["age"]) if data.get("age") not in (None, "") else None
        employee.languages = str(data.get("languages") or "").strip() or None
        employee.certifications = str(data.get("certifications") or "").strip() or None
        employee.technical_experience = str(data.get("technical_experience") or "").strip() or None
        employee.it_years = float(data["it_years"]) if data.get("it_years") not in (None, "") else None
        employee.talent_category = str(data.get("talent_category") or "").strip() or None
        employee.skills = parse_skills(data.get("skills"))
        employee.nationality = normalize_nationality(data.get("nationality"))
        employee_type = normalize_employee_type(data.get("employee_type"))
        if employee_type not in EMPLOYEE_TYPES:
            errors.append(f"第{index}行员工类型无效，已按一般社員处理")
            employee_type = "一般社員"
        employee.employee_type = employee_type or employee.employee_type or "一般社員"

        if is_new:
            db.add(employee)
            db.flush()
            linked_user = ensure_employee_user(db, employee, must_reset_password=True)
            linked_user.role = role_for_employee_type(employee.employee_type)
            created += 1
        else:
            linked_user = ensure_employee_user(db, employee, must_reset_password=False)
            linked_user.role = role_for_employee_type(employee.employee_type)
            updated += 1
        if previous_birth_date != employee.birth_date:
            register_salary_refresh(
                db,
                employee.id,
                months=None,
                include_current=True,
                reason="employee_birth_date_imported",
            )

    db.commit()
    return {"created": created, "updated": updated, "errors": errors}


def import_fixed_resume(data: dict[str, Any], db: Session) -> dict[str, Any]:
    full_name = data.get("full_name")
    if not full_name:
        return {"created": 0, "updated": 0, "errors": ["固定履历书缺少姓名"]}
    email = data.get("email")
    if not email:
        return {"created": 0, "updated": 0, "errors": ["固定履历书缺少邮箱，请在履历书邮件栏补充邮箱后重新导入"]}

    employee = db.scalar(select(Employee).where(Employee.email == email))
    previous_birth_date = employee.birth_date if employee else None
    was_deleted = bool(employee and employee.is_deleted)
    is_new = employee is None or was_deleted
    if is_new:
        if employee:
            employee.is_deleted = False
            employee.deleted_at = None
        else:
            employee = Employee(email=email, full_name=full_name)
    for key, value in data.items():
        if value is not None and hasattr(employee, key):
            setattr(employee, key, value)
    if not employee.phone:
        employee.phone = "未填写"
    if not employee.residence:
        employee.residence = "未填写"
    if not employee.nearest_station:
        employee.nearest_station = "未填写"
    if employee.birth_date and not employee.age:
        today = date.today()
        employee.age = today.year - employee.birth_date.year - ((today.month, today.day) < (employee.birth_date.month, employee.birth_date.day))
    if not employee.graduation_status:
        employee.graduation_status = "未填写"

    if is_new:
        db.add(employee)
        db.flush()
        created, updated = 1, 0
    else:
        created, updated = 0, 1
    employee.employee_type = normalize_employee_type(employee.employee_type)
    ensure_employee_user(db, employee, must_reset_password=is_new)
    if previous_birth_date != employee.birth_date:
        register_salary_refresh(
            db,
            employee.id,
            months=None,
            include_current=True,
            reason="employee_birth_date_imported",
        )
    db.commit()
    return {"created": created, "updated": updated, "errors": []}
